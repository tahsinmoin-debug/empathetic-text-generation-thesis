"""Build version 2 of the human-evaluation study, after the supervisor's revision.

What changed from version 1
---------------------------
1. Rating instead of forced choice. Every reply is now rated 1-5 on three dimensions
   (empathy/emotional validation, contextual relevance, supportiveness) instead of the
   rater picking one winner per situation. A single overall preference question is kept
   at the end of each situation, so the head-to-head win rate reported by the two human
   studies in our reading list is still available.
2. No attention checks. The two decoy replies from the fine-tuned model are gone; those
   situations now show the three real systems like every other item.
3. Fewer situations. Nine ratings per situation instead of one click means the form gets
   long fast: all 13 situations would take about 25 minutes, which is past the point where
   volunteers answer carefully. Eight situations keep it near 15 minutes, and with the
   20-30 raters the supervisor asked for that is 160-240 paired observations per comparison.

Which eight situations
----------------------
Nothing is hand-picked. The set is:
  - the four emotions most central to supportive conversation (sad, lonely, angry, anxious),
    all of them scored by the AI judges, so the human-versus-judge agreement analysis works;
  - the three student situations, because most raters are students;
  - one positive-emotion situation, drawn at random with a fixed seed from the judged
    positive items, so the form is not eight consecutive distressing messages.

Outputs (all in this folder)
----------------------------
  Human_Evaluation_Form_v2.docx     the whole form as text, to read or to copy by hand
  create_google_form_v2.gs          Apps Script that builds the form in Google Drive
  ANSWER_KEY_v2_do_not_share.csv    which response letter is which system
  form_items_v2.csv                 the item set in flat form, for the analysis script

Usage:
    python human_evaluation/build_human_eval_form_v2.py
"""

import json
import random
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PRED_DIR = ROOT / "evaluation_results_d2_settings_comparison"
HERE = Path(__file__).resolve().parent
SEED = 42

OUR_SYSTEM = "two_stage_phi"
BASELINES = ["single_model_same_size", "single_model_bigger"]
SYSTEMS = [OUR_SYSTEM] + BASELINES
DISPLAY = {
    "two_stage_phi": "Our system (two-stage: RoBERTa detector + Phi-4-mini)",
    "single_model_same_size": "Single model, same size (Qwen2.5-1.5B, no detector)",
    "single_model_bigger": "Single model, bigger (Qwen2.5-7B, no detector)",
}

# The conversations carried over from version 1. Kept rather than redrawn so that the
# situations the supervisor has already seen stay the same.
CORE_EMOTIONS = ["sad", "lonely", "angry", "anxious"]          # judged by the AI panel
STUDENT_ITEMS = {                                              # from the wider test set
    "prepared": "hit:6423_conv:12846",
    "furious": "hit:4693_conv:9386",
    "anticipating": "hit:8061_conv:16123",
}

QUESTIONS = [
    ("Empathy / emotional validation",
     "How well does this reply recognise and validate the emotions expressed in the message?",
     "1 = Ignores or dismisses the feelings   3 = Acknowledges them somewhat generically   "
     "5 = Clearly understands and validates the feelings"),
    ("Contextual relevance",
     "How well does this reply address the person's specific situation rather than giving a "
     "generic response?",
     "1 = Generic, could apply to anyone   3 = Partially specific   "
     "5 = Clearly tailored to what was shared"),
    ("Supportiveness",
     "If you had written this message, how supported or comforted would you feel by this reply?",
     "1 = Not at all / worse   3 = Neutral   5 = Very supported"),
]
PREFERENCE_Q = "Overall, which reply responds best to this person's feelings?"

CONSENT = (
    "Thank you for helping with this study.\n\n"
    "You will read 8 short messages from people describing something they are going through. "
    "For each message you will see three possible replies, and you will rate each reply on "
    "three simple 1-5 scales. There are no right or wrong answers - we want your honest "
    "impression.\n\n"
    "This takes about 15 minutes. Your answers are completely anonymous: we do not collect "
    "your name or email address. Some messages describe sad or stressful situations, and you "
    "may stop at any time.\n\n"
    "The replies were written by different computer programs, but we do not say which is "
    "which, because that would change how you read them.\n\n"
    "By continuing, you agree to take part."
)


# Typography is normalised so the three replies cannot be told apart by their punctuation:
# some systems produce curly quotes and dashes, others straight ones, and a careful rater
# could otherwise notice that pattern across situations. Only punctuation is changed.
TYPOGRAPHY = {"’": "'", "‘": "'", "“": '"', "”": '"',
              "—": " - ", "–": "-", "…": "..."}


def clean(text):
    out = str(text).replace("_comma_", ",")
    for src, dst in TYPOGRAPHY.items():
        out = out.replace(src, dst)
    return " ".join(out.split())


def load_systems():
    return {name: pd.read_csv(PRED_DIR / f"{name}_scored_predictions.csv").set_index("conv_id")
            for name in SYSTEMS}


def select_items(systems):
    key = pd.read_csv(HERE / "ANSWER_KEY_do_not_share_with_raters.csv")
    v1 = key.groupby("item_number").agg(conv_id=("conv_id", "first"),
                                        emotion=("gold_emotion", "first"),
                                        judged=("ai_judge_scored", "first"))
    rng = random.Random(SEED)

    chosen = []
    for emotion in CORE_EMOTIONS:
        row = v1[v1["emotion"] == emotion].iloc[0]
        chosen.append((emotion, row["conv_id"], bool(row["judged"])))
    for emotion, conv_id in STUDENT_ITEMS.items():
        chosen.append((emotion, conv_id, False))

    # one positive-emotion situation, drawn at random from the judged items of version 1
    positive = [(r["emotion"], r["conv_id"]) for _, r in v1.iterrows()
                if r["judged"] and r["emotion"] in
                {"joyful", "content", "impressed", "nostalgic", "faithful", "proud", "grateful"}]
    emotion, conv_id = rng.choice(sorted(positive))
    chosen.append((emotion, conv_id, True))

    rng.shuffle(chosen)

    # Counterbalanced ordering: the three cyclic rotations put each system in each position
    # once, so repeating them across eight items keeps position bias close to balanced
    # (3/3/2 per position). Which rotation lands on which item is randomised.
    rotations = [(0, 1, 2), (1, 2, 0), (2, 0, 1)]
    orders = [rotations[i % 3] for i in range(len(chosen))]
    rng.shuffle(orders)

    items = []
    for n, ((emotion, conv_id, judged), rotation) in enumerate(zip(chosen, orders), start=1):
        ordered = [SYSTEMS[i] for i in rotation]
        items.append({
            "number": n,
            "conv_id": conv_id,
            "gold_emotion": emotion,
            "ai_judge_scored": judged,
            "user_text": clean(systems[OUR_SYSTEM].loc[conv_id, "user_text"]),
            "replies": [{"label": "ABC"[pos],
                         "system": name,
                         "text": clean(systems[name].loc[conv_id, "generated_response"]),
                         "words": int(systems[name].loc[conv_id, "word_count"])}
                        for pos, name in enumerate(ordered)],
        })
    return items


# ------------------------------------------------------------------ outputs

def write_answer_key(items):
    rows = []
    for it in items:
        for r in it["replies"]:
            rows.append({"item_number": it["number"], "conv_id": it["conv_id"],
                         "gold_emotion": it["gold_emotion"],
                         "ai_judge_scored": it["ai_judge_scored"],
                         "shown_as": f"Response {r['label']}",
                         "actual_system": r["system"],
                         "system_description": DISPLAY[r["system"]],
                         "word_count": r["words"]})
    df = pd.DataFrame(rows)
    df.to_csv(HERE / "ANSWER_KEY_v2_do_not_share.csv", index=False)
    df.to_csv(HERE / "form_items_v2.csv", index=False)
    return df


def write_apps_script(items):
    """Apps Script that builds the whole form, so nobody types 80 questions by hand."""
    payload = json.dumps([{ "n": it["number"], "msg": it["user_text"],
                            "replies": [{"label": r["label"], "text": r["text"]}
                                        for r in it["replies"]]} for it in items],
                         ensure_ascii=True, indent=2)
    questions = json.dumps([{"title": q[1], "help": q[2]} for q in QUESTIONS],
                           ensure_ascii=True, indent=2)
    script = f'''/**
 * Creates the human-evaluation Google Form automatically.
 *
 * How to run it:
 *   1. Go to script.google.com and choose "New project".
 *   2. Delete whatever is in the editor, paste this whole file, and press Save.
 *   3. Press Run. Google asks for permission the first time - allow it.
 *   4. The execution log prints the form's edit link and its public link.
 *
 * Nothing in the form names a model. Raters see only Response A, B and C.
 */

var ITEMS = {payload};

var QUESTIONS = {questions};

var PREFERENCE_Q = {json.dumps(PREFERENCE_Q)};

var CONSENT = {json.dumps(CONSENT)};

function createForm() {{
  var form = FormApp.create('How do people respond to someone who is upset? (15 minutes)');
  form.setDescription(CONSENT);
  form.setProgressBar(true);
  try {{ form.setCollectEmail(false); }} catch (e) {{ /* newer accounts already collect no email */ }}
  form.setShuffleQuestions(false);

  form.addMultipleChoiceItem()
      .setTitle('I am 18 or older and I agree to take part in this study.')
      .setChoiceValues(['Yes', 'No'])
      .setRequired(true);

  for (var i = 0; i < ITEMS.length; i++) {{
    var item = ITEMS[i];
    form.addPageBreakItem().setTitle('Situation ' + item.n + ' of ' + ITEMS.length);

    var labels = [];
    form.addSectionHeaderItem()
        .setTitle('Someone wrote:')
        .setHelpText('"' + item.msg + '"');
    for (var r = 0; r < item.replies.length; r++) {{
      form.addSectionHeaderItem()
          .setTitle('Response ' + item.replies[r].label)
          .setHelpText(item.replies[r].text);
      labels.push('Response ' + item.replies[r].label);
    }}

    for (var q = 0; q < QUESTIONS.length; q++) {{
      form.addGridItem()
          .setTitle(QUESTIONS[q].title)
          .setHelpText(QUESTIONS[q].help)
          .setRows(labels)
          .setColumns(['1', '2', '3', '4', '5'])
          .setRequired(true);
    }}

    form.addMultipleChoiceItem()
        .setTitle(PREFERENCE_Q)
        .setChoiceValues(labels)
        .setRequired(true);
  }}

  form.addPageBreakItem().setTitle('Thank you');
  form.addParagraphTextItem()
      .setTitle('Anything you would like to add? (optional)')
      .setRequired(false);

  Logger.log('EDIT THE FORM HERE:   ' + form.getEditUrl());
  Logger.log('SEND THIS TO RATERS:  ' + form.getPublishedUrl());
  return form.getPublishedUrl();
}}

/**
 * Apps Script runs whichever function is selected in the toolbar, and a new project
 * starts with an empty myFunction. If that one is still selected, running it would do
 * nothing and print no links, so it simply calls createForm.
 */
function myFunction() {{
  createForm();
}}
'''
    (HERE / "create_google_form_v2.gs").write_text(script, encoding="utf-8")


def write_document(items, key):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    def h(t, lvl):
        doc.add_heading(t, level=lvl)

    def p(t, bold=False, italic=False, size=11, color=None):
        para = doc.add_paragraph()
        run = para.add_run(t)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        para.paragraph_format.space_after = Pt(6)

    h("Human Evaluation Study - Version 2", 0)
    p("Context-aware empathetic text generation: rating study", italic=True, size=12)
    p("Every word the raters see, plus the build instructions and the analysis plan.", size=10)

    h("1. What changed and why", 1)
    p("Three replies per situation, as before. What changed is how they are judged.")
    for line in [
        "Each reply is now rated 1-5 on three dimensions instead of the rater choosing one "
        "winner. This gives a score for every system on every situation, so we can measure "
        "how much better one system is, not only how often it wins.",
        "The two attention-check situations no longer contain a deliberately poor reply. "
        "Those situations now show the three real systems like every other item.",
        "Eight situations instead of thirteen. Nine ratings per situation means all thirteen "
        "would take about 25 minutes; past roughly 15 minutes volunteers stop reading "
        "carefully, and careless answers add noise rather than evidence.",
        "Target: 20 to 30 people who answer seriously. That gives 160 to 240 paired "
        "observations for each comparison, since every rater rates all three systems on all "
        "eight situations.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    h("2. Which situations are in the form", 1)
    p("Nothing is hand-picked. The eight are: the four emotions most central to supportive "
      "conversation (sad, lonely, angry, anxious), which the AI judges also scored, so human "
      "and AI ratings can be compared; the three student situations, because most raters are "
      "students; and one positive-emotion situation drawn at random with a fixed seed, so the "
      "form is not eight distressing messages in a row.")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for cell, head in zip(t.rows[0].cells, ["Situation", "Emotion", "Also scored by the AI judges", "Why it is here"]):
        cell.text = head
    for it in items:
        cells = t.add_row().cells
        why = ("core supportive emotion" if it["gold_emotion"] in CORE_EMOTIONS
               else "student situation" if it["gold_emotion"] in STUDENT_ITEMS
               else "positive emotion, drawn at random")
        for cell, v in zip(cells, [it["number"], it["gold_emotion"],
                                   "yes" if it["ai_judge_scored"] else "no", why]):
            cell.text = str(v)
    doc.add_paragraph()

    h("3. How to build the form", 1)
    p("The fastest way is the script: open script.google.com, create a new project, paste "
      "create_google_form_v2.gs, press Run, and Google builds all 80 questions. The log prints "
      "the link to send out. If you would rather build it by hand, Section 5 has every word.",)
    p("Settings that matter:", bold=True)
    for line in ["Do not collect email addresses - the study is anonymous.",
                 "Leave 'Shuffle question order' off. The reply order is already balanced "
                 "across situations; shuffling on top of that breaks the answer key.",
                 "Mark every rating required, so there are no half-finished responses.",
                 "Never write a model's name in the form. Raters see only Response A, B and C."]:
        doc.add_paragraph(line, style="List Bullet")

    h("4. The three questions asked about every reply", 1)
    for title, question, anchors in QUESTIONS:
        p(title, bold=True)
        p(question)
        p(anchors, size=10, italic=True)
    p(PREFERENCE_Q, bold=True)
    p("Asked once per situation, after the three ratings.", size=10, italic=True)

    h("5. The complete item set", 1)
    p("This is exactly what a rater sees. The system names in brackets are for you only and "
      "must not appear in the form.", size=10, italic=True)
    for it in items:
        h(f"Situation {it['number']} of {len(items)}", 2)
        p(f"Someone wrote: “{it['user_text']}”")
        for r in it["replies"]:
            p(f"Response {r['label']}:", bold=True)
            p(r["text"])
            p(f"[{DISPLAY[r['system']]}, {r['words']} words - not shown to raters]",
              size=9, italic=True, color=RGBColor(0x88, 0x88, 0x88))
        p("Then: the three 1-5 questions about each reply, and the preference question.",
          size=10, italic=True)

    h("6. How the answers will be analysed", 1)
    for line in [
        "Ratings are paired: every rater rates all three systems on the same situation, so the "
        "comparisons use the Wilcoxon signed-rank test on per-rater mean scores, with Holm "
        "correction across the three dimensions and the three system pairs.",
        "The preference question gives a win rate per system, tested against the 33.3% chance "
        "level with a binomial test.",
        "Agreement between raters is reported with the intraclass correlation for the 1-5 "
        "scores and Fleiss' kappa for the preference question.",
        "Agreement with the AI judges is the point of the study: the human mean score per "
        "system is correlated with the AI-judge mean on the five situations the judges also "
        "scored, and the two rankings are compared.",
        "Anyone who answers 'No' to the consent question, or who gives the same rating to "
        "every reply in every situation, is excluded before analysis. These are the only two "
        "exclusion rules, and both are fixed in advance.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    p("Limitation to state in the thesis: three of the eight situations are not scored by the "
      "AI judges, so the human-versus-judge agreement analysis runs on the other five.",
      size=10, italic=True)

    doc.save(HERE / "Human_Evaluation_Form_v2.docx")


def main():
    systems = load_systems()
    items = select_items(systems)
    key = write_answer_key(items)
    write_apps_script(items)
    write_document(items, key)

    ratings = len(items) * len(SYSTEMS) * len(QUESTIONS)
    print(f"{len(items)} situations, {ratings} ratings + {len(items)} preference questions")
    for it in items:
        order = " ".join(f"{r['label']}={r['system']}" for r in it["replies"])
        print(f"  {it['number']}. {it['gold_emotion']:13s} judged={str(it['ai_judge_scored']):5s} {order}")
    counts = {}
    for it in items:
        for r in it["replies"]:
            counts.setdefault(r["system"], [0, 0, 0])["ABC".index(r["label"])] += 1
    print("position balance (A/B/C):")
    for name, c in counts.items():
        print(f"  {name:26s} {c}")


if __name__ == "__main__":
    main()
