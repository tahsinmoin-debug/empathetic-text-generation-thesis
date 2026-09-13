"""Convert every WRITING_GUIDE.md (and the kit README) into a Word document next to it.

Markers in the drafts are expanded in the Word version:
    [Table: name.csv]    -> the actual table from tables/name.csv
    [Figure: name.png]   -> the actual image from figures/name.png
so teammates get a document they can edit directly, without opening CSVs or PNGs.

Run after build_writing_tables.py and build_writing_figures.py.

Usage:
    python writing/_scripts/build_writing_docx.py
"""

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

W = Path(__file__).resolve().parents[1]
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[\[[^\]]+\]\])")
MARKER = re.compile(r"\[(Table|Figure): ([^\]]+?\.(?:csv|png))\]")
PLACEHOLDER_COLOR = RGBColor(0xB0, 0x3A, 0x00)


def add_runs(paragraph, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("[[") and part.endswith("]]"):
            run = paragraph.add_run(part)
            run.bold = True
            run.font.color.rgb = PLACEHOLDER_COLOR
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc, header, rows):
    ncols = len(header)
    font = 8 if ncols <= 6 else 6.5
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        add_runs(cell.paragraphs[0], str(text), size=font)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shade(cell, "E3EEFB")
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            add_runs(cell.paragraphs[0], str(text), size=font)
    doc.add_paragraph()


def csv_table(doc, path):
    df = pd.read_csv(path)
    df = df.fillna("")
    rows = []
    for _, r in df.iterrows():
        rows.append([f"{v:.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else v for v in r.tolist()])
    add_table(doc, list(df.columns), rows)


def expand_markers(doc, folder, line):
    """Insert real tables/figures for every marker; return the text left over."""
    for kind, name in MARKER.findall(line):
        if kind == "Table":
            path = folder / "tables" / name
            if path.exists():
                cap = doc.add_paragraph()
                add_runs(cap, f"*Table: {name}*", size=9)
                csv_table(doc, path)
            else:
                add_runs(doc.add_paragraph(), f"[[MISSING TABLE: {name}]]")
        else:
            path = folder / "figures" / name
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.2))
                cap = doc.add_paragraph()
                add_runs(cap, f"*Figure: {name}*", size=9)
            else:
                add_runs(doc.add_paragraph(), f"[[MISSING FIGURE: {name}]]")
    return MARKER.sub("", line).strip()


def convert(md_path, folder):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, header, rows)
            continue

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        if MARKER.search(stripped):
            rest = expand_markers(doc, folder, stripped)
            if rest:
                add_runs(doc.add_paragraph(), rest)
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1)) - 1
            h = doc.add_heading(level=level)
            add_runs(h, heading.group(2))
        elif stripped.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, stripped[2:])
        elif re.match(r"^- \[ \]\s", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, "☐ " + stripped[6:])
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", stripped))
        else:
            add_runs(doc.add_paragraph(), stripped)
        i += 1

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def main():
    made = []
    for md in sorted(W.glob("*/WRITING_GUIDE.md")):
        made.append(convert(md, md.parent))
    made.append(convert(W / "README.md", W))
    for m in made:
        print("wrote", m.relative_to(W))


if __name__ == "__main__":
    main()
