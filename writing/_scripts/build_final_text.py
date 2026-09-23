"""Build FINAL_TEXT.md and FINAL_TEXT.docx in every section folder, plus writing/REFERENCES_ALL.

Source texts live in writing/_scripts/final_text_src/. They use three markers:
    {cite:key1,key2}                         -> IEEE numbered citation, e.g. [4], [7]
    [[TABLE 5.3 | file.csv | Caption text]]  -> the table from that section's tables/ folder
    [[FIGURE 5.1 | file.png | Caption text]] -> the image from that section's figures/ folder
Citation numbers follow order of first appearance across the thesis (Chapter 2 first), so a paper
keeps the same number in every section. Build fails loudly on an unknown citation key, a missing
table/figure file, or a table/figure number used twice.

Captions go BELOW tables and figures, as in the BRAC CSE400 template.
[[TEAM TO FILL: ...]] and [[PENDING: ...]] placeholders are kept and highlighted in the Word output.

Usage:
    python writing/_scripts/build_final_text.py
"""

import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
W = HERE.parent
SRC = HERE / "final_text_src"
sys.path.insert(0, str(HERE))
from references import REFS  # noqa: E402
from build_writing_docx import add_runs, add_table  # noqa: E402

# Thesis order - decides citation numbering.
ORDER = [
    ("ch2_1", "Chapter 2 - 2.1 Preliminaries"),
    ("ch2_2", "Chapter 2 - 2.2 Review of Existing Research"),
    ("ch2_3", "Chapter 2 - 2.3 Summary of Key Findings"),
    ("ch3_4", "Chapter 3 - 3.4 Ethical Issues"),
    ("ch3_6", "Chapter 3 - 3.6 Project Management Plan"),
    ("ch3_7", "Chapter 3 - 3.7 Risk Management"),
    ("ch3_8", "Chapter 3 - 3.8 Economic Analysis"),
    ("ch4_1", "Chapter 4 - 4.1 Design Process or Methodology Overview"),
    ("ch4_2", "Chapter 4 - 4.2 Preliminary Design or Design (Model) Specification"),
    ("ch5_1", "Chapter 5 - 5.1 Performance Evaluation"),
    ("ch5_2a", "Chapter 5 - 5.2 Analysis of Design Solutions"),
    ("ch5_2b", "Chapter 5 - 5.2 Final Design Adjustments"),
    ("ch5_3", "Chapter 5 - 5.3 Statistical Analysis"),
    ("ch5_4", "Chapter 5 - 5.4 Comparisons and Relationships"),
    ("ch5_6", "Chapter 5 - 5.6 Discussion"),
    ("tools", "Presentation - Tools Used"),
]

CITE = re.compile(r"\{cite:([a-z0-9_,\s]+)\}")
TABLE = re.compile(r"\[\[TABLE ([0-9.]+) \| ([^|\]]+?) \| ([^\]]+?)\]\]")
FIGURE = re.compile(r"\[\[FIGURE ([0-9.]+) \| ([^|\]]+?) \| ([^\]]+?)\]\]")


def number_citations(texts):
    numbers = {}
    for _, text in texts:
        for group in CITE.findall(text):
            for key in [k.strip() for k in group.split(",") if k.strip()]:
                if key not in REFS:
                    raise SystemExit(f"Unknown citation key: {key}")
                numbers.setdefault(key, len(numbers) + 1)
    return numbers


def resolve_cites(text, numbers):
    def repl(m):
        keys = [k.strip() for k in m.group(1).split(",") if k.strip()]
        return ", ".join(f"[{numbers[k]}]" for k in keys)
    return CITE.sub(repl, text)


def fmt(v):
    """Print a cell as the CSV wrote it, and only shorten a number left at full precision.

    Reading the file as text is what keeps "8.00" from collapsing to "8", "+0.59" from losing
    its sign, and a p-value of 8e-05 from printing as "0", all of which happened when pandas
    was allowed to parse the numbers first.
    """
    s = "" if v is None else str(v).strip()
    if s in ("", "nan", "None"):
        return ""
    long_float = re.fullmatch(r"[-+]?\d+\.(\d{4,})", s)
    return f"{float(s):.2f}" if long_float else s


def csv_as_markdown(path):
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    lines = ["| " + " | ".join(df.columns) + " |", "|" + "|".join("---" for _ in df.columns) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(fmt(v).replace("|", "/") for v in r.tolist()) + " |")
    return "\n".join(lines)


def check_numbers(texts):
    seen = {}
    for key, text in texts:
        if key == "tools":  # deliberately repeats Table 4.3 for the presentation
            continue
        for kind, rx in (("Table", TABLE), ("Figure", FIGURE)):
            for num, _, _ in rx.findall(text):
                tag = f"{kind} {num}"
                if tag in seen:
                    raise SystemExit(f"{tag} used twice: in {seen[tag]} and {key}")
                seen[tag] = key


def to_markdown(text, folder, numbers, used):
    out = []
    for line in text.splitlines():
        t = TABLE.fullmatch(line.strip())
        f = FIGURE.fullmatch(line.strip())
        if t:
            num, fname, cap = t.groups()
            path = folder / "tables" / fname.strip()
            if not path.exists():
                raise SystemExit(f"Missing table file {path}")
            out += ["", csv_as_markdown(path), "", f"*Table {num}: {cap.strip()}*", ""]
        elif f:
            num, fname, cap = f.groups()
            path = folder / "figures" / fname.strip()
            if not path.exists():
                raise SystemExit(f"Missing figure file {path}")
            out += ["", f"![Figure {num}](figures/{fname.strip()})", "", f"*Figure {num}: {cap.strip()}*", ""]
        else:
            out.append(resolve_cites(line, numbers))
    out += ["", "## References", ""]
    out += [f"[{numbers[k]}] {REFS[k]}  " for k in sorted(used, key=lambda k: numbers[k])]
    return "\n".join(out) + "\n"


def to_docx(md_text, folder, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in s.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, header, rows)
            continue
        img = re.match(r"^!\[[^\]]*\]\((figures/[^)]+)\)$", s)
        if img:
            doc.add_picture(str(folder / img.group(1)), width=Inches(6.0))
        elif not s:
            pass
        elif s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            h = doc.add_heading(level=min(level - 1, 3))
            add_runs(h, s.lstrip("#").strip())
        elif s.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, s[2:])
        elif s.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\*(Table|Figure) [0-9.]+:.*\*$", s):
            p = doc.add_paragraph()
            add_runs(p, s, size=10)
        elif re.match(r"^\[\d+\] ", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.4)
            add_runs(p, s.rstrip(), size=10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, s)
        i += 1
    doc.save(out_path)


def main():
    texts = []
    for key, folder in ORDER:
        src = SRC / f"{key}.md"
        if not src.exists():
            raise SystemExit(f"Missing source text {src}")
        texts.append((key, src.read_text(encoding="utf-8")))
    check_numbers(texts)
    numbers = number_citations(texts)

    for (key, folder), (_, text) in zip(ORDER, texts):
        used = {k.strip() for g in CITE.findall(text) for k in g.split(",") if k.strip()}
        md = to_markdown(text, W / folder, numbers, used)
        (W / folder / "FINAL_TEXT.md").write_text(md, encoding="utf-8")
        to_docx(md, W / folder, W / folder / "FINAL_TEXT.docx")
        print(f"wrote {folder}/FINAL_TEXT  ({len(used)} citations)")

    all_md = ["# References (whole thesis, in order of first citation)", "",
              "> Numbering assumes Chapter 1 cites nothing new. If Chapter 1 adds citations, add them to "
              "`writing/_scripts/references.py`, cite them in a Chapter 1 source file, and rebuild.", ""]
    all_md += [f"[{n}] {REFS[k]}  " for k, n in sorted(numbers.items(), key=lambda kv: kv[1])]
    unused = sorted(set(REFS) - set(numbers))
    if unused:
        print("note - references defined but not cited anywhere:", ", ".join(unused))
    text = "\n".join(all_md) + "\n"
    (W / "REFERENCES_ALL.md").write_text(text, encoding="utf-8")
    to_docx(text, W, W / "REFERENCES_ALL.docx")
    print(f"wrote REFERENCES_ALL ({len(numbers)} references)")


if __name__ == "__main__":
    main()
