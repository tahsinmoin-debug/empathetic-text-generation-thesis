"""Build the ICCIT (IEEE conference) paper from paper_src.md.

Outputs, next to this folder's parent:
    paper.tex           IEEEtran conference LaTeX - upload with figures/ to Overleaf
    paper.docx          Word version in IEEE two-column layout, for editing
    PAPER_PREVIEW.md    plain preview (citations and table numbers resolved)

Source markers:
    {cite:key1, key2}                      -> IEEE numbered citations, numbered by first appearance
    {ref:tab:main} / {ref:fig:arch}        -> table (Roman) / figure (Arabic) number
    [[TABLE label | caption | wide|narrow]] ... pipe table ... [[/TABLE]]
    [[FIGURE label | file.png | caption]]
    [[PENDING: ...]]                       -> kept and highlighted for the authors

Reference texts are shared with the thesis kit (writing/_scripts/references.py) so both use
identical, verified entries; paper-only references are in EXTRA_REFS.

Usage:
    python "ICCIT format writing/_build/build_paper.py"
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
ROOT = OUT.parent
sys.path.insert(0, str(ROOT / "writing" / "_scripts"))
from references import REFS  # noqa: E402

EXTRA_REFS = {
    "yang2024hef": 'Z. Yang, Z. Ren, Y. Wang, S. Peng, H. Sun, X. Zhu, and X. Liao, "Enhancing empathetic response '
                   'generation by augmenting LLMs with small-scale empathetic models," arXiv:2402.11801, 2024.',
}
ALL_REFS = {**REFS, **EXTRA_REFS}

CITE = re.compile(r"\{cite:([a-z0-9_,\s]+)\}")
REF = re.compile(r"\{ref:((?:tab|fig):[a-z]+)\}")
PENDING = re.compile(r"\[\[PENDING:[^\]]*\]\]")
TABLE_OPEN = re.compile(r"^\[\[TABLE ([a-z:]+) \| (.+?) \| (wide|narrow)\]\]$")
FIGURE = re.compile(r"^\[\[FIGURE ([a-z:]+) \| ([^|]+?) \| (.+?)\]\]$")
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


# ------------------------------------------------------------------ parsing

def parse(text):
    """Return title and a list of blocks: (kind, payload)."""
    lines = text.splitlines()
    title = lines[0].removeprefix("TITLE:").strip()
    blocks, para, i = [], [], 1

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    while i < len(lines):
        s = lines[i].strip()
        t = TABLE_OPEN.match(s)
        f = FIGURE.match(s)
        if not s:
            flush()
        elif s.startswith("### "):
            flush(); blocks.append(("h2", re.sub(r"^[A-Z]\.\s+", "", s[4:])))
        elif s.startswith("## "):
            flush(); blocks.append(("h1", re.sub(r"^[IVX]+\.\s+", "", s[3:])))
        elif s.startswith("- "):
            flush()
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:]); i += 1
            blocks.append(("list", items)); continue
        elif t:
            flush()
            rows = []
            i += 1
            while not lines[i].strip().startswith("[[/TABLE]]"):
                r = lines[i].strip()
                if r and not re.match(r"^\|[\s:|-]+\|$", r):
                    rows.append([c.strip() for c in r.strip("|").split("|")])
                i += 1
            blocks.append(("table", {"label": t.group(1), "caption": t.group(2), "wide": t.group(3) == "wide",
                                     "rows": rows}))
        elif f:
            flush(); blocks.append(("figure", {"label": f.group(1), "file": f.group(2), "caption": f.group(3)}))
        else:
            para.append(s)
        i += 1
    flush()
    return title, blocks


def all_text(blocks):
    for kind, pl in blocks:
        if kind in ("p", "h1", "h2"):
            yield pl
        elif kind == "list":
            yield from pl
        elif kind == "table":
            yield pl["caption"]
            for r in pl["rows"]:
                yield from r
        elif kind == "figure":
            yield pl["caption"]


def numbering(blocks):
    cites = {}
    for s in all_text(blocks):
        for g in CITE.findall(s):
            for k in [k.strip() for k in g.split(",") if k.strip()]:
                if k not in ALL_REFS:
                    raise SystemExit(f"Unknown citation key: {k}")
                cites.setdefault(k, len(cites) + 1)
    refs, nt, nf = {}, 0, 0
    for kind, pl in blocks:
        if kind == "table":
            refs[pl["label"]] = ROMAN[nt]; nt += 1
        elif kind == "figure":
            nf += 1; refs[pl["label"]] = str(nf)
    for s in all_text(blocks):
        for lab in REF.findall(s):
            if lab not in refs:
                raise SystemExit(f"Unknown reference label: {lab}")
    return cites, refs


# ------------------------------------------------------------------ LaTeX

def tex_inline(s, cites):
    keep = {}

    def stash(val):
        key = f"@@{len(keep)}@@"
        keep[key] = val
        return key

    s = PENDING.sub(lambda m: stash(r"\textcolor{red}{" + tex_escape(m.group(0)[2:-2]) + "}"), s)
    s = CITE.sub(lambda m: stash(r"\cite{" + ",".join(cites.get(k.strip(), k.strip()) if isinstance(cites, dict) else k.strip()
                                                       for k in m.group(1).split(",")) + "}"), s)
    s = REF.sub(lambda m: stash(r"\ref{" + m.group(1) + "}"), s)
    s = tex_escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", s)
    s = re.sub(r"\*(.+?)\*", r"\\emph{\1}", s)
    s = re.sub(r"\brho\b", r"$\\rho$", s)
    s = s.replace("Fig. @@", "Fig.~@@").replace("Table @@", "Table~@@").replace("Section V-C", "Section~V-C")
    s = tex_quotes(s)
    for k, v in keep.items():
        s = s.replace(k, v)
    return s


def tex_escape(s):
    s = s.replace("\\", r"\textbackslash{}")
    for a, b in [("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"), ("_", r"\_"),
                 ("{", r"\{"), ("}", r"\}"), ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}"),
                 ("<", r"$<$"), (">", r"$>$")]:
        s = s.replace(a, b)
    return s.replace(r"\textbackslash\{\}", r"\textbackslash{}")


def tex_quotes(s):
    """Straight quotes render as two closing marks in LaTeX; convert pairs to ``...'' form."""
    out, opening = [], True
    for ch in s:
        if ch == '"':
            out.append("``" if opening else "''")
            opening = not opening
        else:
            out.append(ch)
    return "".join(out)


def tex_bib(entry):
    """Template style for references: LaTeX quotes and en-dashes in page ranges."""
    entry = tex_quotes(tex_escape(entry))
    return re.sub(r"(pp\.\s*\d+)-(\d+)", "\\1--\\2", entry)


def to_tex(title, blocks, cites):
    """Emit LaTeX in the official IEEE conference template's structure.

    Preamble, author block, table style (\hline, caption above, centred) and the b1..bN bibliography
    keys follow IEEE-conference-template-062824.tex, so the file can be dropped into Overleaf or
    Prism and compiled without further changes. All template guidance text is removed, as IEEE
    requires.
    """
    key_of = {k: f"b{n}" for k, n in cites.items()}  # template-style citation labels
    out = [
        r"\documentclass[conference]{IEEEtran}",
        r"\IEEEoverridecommandlockouts",
        r"% Template version as of 6/27/2024",
        "",
        r"\usepackage{cite}",
        r"\usepackage{amsmath,amssymb,amsfonts}",
        r"\usepackage{algorithmic}",
        r"\usepackage{graphicx}",
        r"\usepackage{textcomp}",
        r"\usepackage{xcolor}",
        r"\def\BibTeX{{\rm B\kern-.05em{\sc i\kern-.025em b}\kern-.08em",
        r"    T\kern-.1667em\lower.7ex\hbox{E}\kern-.125emX}}",
        "",
        r"\begin{document}",
        "",
        r"\title{" + tex_escape(title) + "}",
        "",
        r"% DOUBLE-BLIND SUBMISSION: no names, affiliations or e-mail addresses (ICCIT requirement).",
        r"% For the camera-ready version, replace this block with the author blocks of the template:",
        r"%   \author{\IEEEauthorblockN{1\textsuperscript{st} Given Name Surname}",
        r"%   \IEEEauthorblockA{\textit{dept. name of organization} \\ \textit{name of organization} \\",
        r"%   City, Country \\ email address or ORCID}",
        r"%   \and ... }",
        r"\author{\IEEEauthorblockN{Anonymous Authors}}",
        "",
        r"\maketitle",
        "",
    ]
    mode = None
    for kind, pl in blocks:
        if kind == "h1":
            if mode == "abstract":
                out += [r"\end{abstract}", ""]
            elif mode == "keywords":
                out += [r"\end{IEEEkeywords}", ""]
            mode = None
            if pl == "Abstract":
                out.append(r"\begin{abstract}"); mode = "abstract"
            elif pl == "Index Terms":
                out.append(r"\begin{IEEEkeywords}"); mode = "keywords"
            elif pl == "Acknowledgment":
                out += ["", r"\section*{Acknowledgment}"]
            else:
                out += ["", r"\section{" + tex_inline(pl, key_of) + "}"]
        elif kind == "h2":
            out += ["", r"\subsection{" + tex_inline(pl, key_of) + "}"]
        elif kind == "p":
            out += [tex_inline(pl, key_of), ""]
        elif kind == "list":
            out.append(r"\begin{itemize}")
            out += [r"\item " + tex_inline(x, key_of) for x in pl]
            out += [r"\end{itemize}", ""]
        elif kind == "figure":
            out += [r"\begin{figure}[htbp]",
                    r"\centerline{\includegraphics[width=\columnwidth]{figures/" + pl["file"] + "}}",
                    r"\caption{" + tex_inline(pl["caption"], key_of) + "}",
                    r"\label{" + pl["label"] + "}", r"\end{figure}", ""]
        elif kind == "table":
            env = "table*" if pl["wide"] else "table"
            ncol = len(pl["rows"][0])
            # Tables with many columns are scaled to the text/column width so they cannot overflow;
            # narrow ones are simply set smaller, which keeps the template's look.
            fit = ncol >= 5
            width = r"\textwidth" if pl["wide"] else r"\columnwidth"
            out += [r"\begin{" + env + "}[htbp]",
                    r"\caption{" + tex_inline(pl["caption"], key_of) + "}",
                    r"\label{" + pl["label"] + "}",
                    r"\begin{center}"]
            out.append(r"\resizebox{" + width + r"}{!}{%" if fit else r"\scriptsize")
            out += [r"\begin{tabular}{|l|" + "c|" * (ncol - 1) + "}", r"\hline"]
            for j, row in enumerate(pl["rows"]):
                cells = [tex_inline(c, key_of) for c in row]
                if j == 0:
                    cells = [r"\textbf{" + c + "}" for c in cells]
                out += [" & ".join(cells) + r" \\", r"\hline"]
            out += [r"\end{tabular}" + ("}" if fit else ""), r"\end{center}", r"\end{" + env + "}", ""]
    if mode == "abstract":
        out.append(r"\end{abstract}")
    elif mode == "keywords":
        out.append(r"\end{IEEEkeywords}")

    out += ["", r"\begin{thebibliography}{00}"]
    for k, n in sorted(cites.items(), key=lambda kv: kv[1]):
        out.append(r"\bibitem{" + key_of[k] + "} " + tex_bib(ALL_REFS[k]))
    out += [r"\end{thebibliography}", "", r"\end{document}", ""]
    return "\n".join(out)


# ------------------------------------------------------------------ Word

def cite_text(m, cites):
    return ", ".join(f"[{cites[k.strip()]}]" for k in m.group(1).split(",") if k.strip())


def plain(s, cites, refs):
    s = CITE.sub(lambda m: cite_text(m, cites), s)
    s = REF.sub(lambda m: refs[m.group(1)], s)
    return re.sub(r"\brho\b", "\u03c1", s)


def runs(par, s, cites, refs, size=10, bold=False, italic=False):
    s = plain(s, cites, refs)
    for piece in re.split(r"(\[\[PENDING:[^\]]*\]\]|\*\*.+?\*\*|\*[^*]+?\*)", s):
        if not piece:
            continue
        r = par.add_run()
        if piece.startswith("[[PENDING"):
            r.text = piece[2:-2]
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif piece.startswith("**"):
            r.text, r.bold = piece[2:-2], True
        elif piece.startswith("*"):
            r.text, r.italic = piece[1:-1], True
        else:
            r.text = piece
        r.font.size = Pt(size)
        r.font.name = "Times New Roman"
        if bold:
            r.bold = True
        if italic:
            r.italic = True


def set_columns(section, n):
    sp = section._sectPr
    for c in sp.findall(qn("w:cols")):
        sp.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")
    sp.append(cols)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, indent=0.0):
    pf = par.paragraph_format
    par.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.first_line_indent = Inches(indent)
    pf.line_spacing = 1.0


def to_docx(title, blocks, cites, refs, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
    sec.left_margin = sec.right_margin = Inches(0.625)

    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    runs(p, title, cites, refs, size=22)
    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=12)
    runs(p, "Anonymous Authors (submitted for double-blind review)", cites, refs, size=10, italic=True)
    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body, 2)

    mode, sec_no, sub_no = None, 0, 0
    for kind, pl in blocks:
        if kind == "h1":
            mode = pl if pl in ("Abstract", "Index Terms") else None
            if mode:
                continue
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4)
            if pl == "Acknowledgment":
                runs(p, "ACKNOWLEDGMENT", cites, refs, size=10)
            else:
                sub_no = 0
                runs(p, f"{ROMAN[sec_no]}. {pl.upper()}", cites, refs, size=10)
                sec_no += 1
        elif kind == "h2":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3)
            runs(p, f"{chr(65 + sub_no)}. {pl}", cites, refs, size=10, italic=True)
            sub_no += 1
        elif kind == "p":
            p = doc.add_paragraph()
            if mode == "Abstract":
                fmt(p, after=4)
                runs(p, "Abstract\u2014", cites, refs, size=9, bold=True, italic=True)
                runs(p, pl, cites, refs, size=9, bold=True)
            elif mode == "Index Terms":
                fmt(p, after=8)
                runs(p, "Index Terms\u2014", cites, refs, size=9, bold=True, italic=True)
                runs(p, pl, cites, refs, size=9, bold=True)
            else:
                fmt(p, indent=0.125)
                runs(p, pl, cites, refs)
        elif kind == "list":
            for x in pl:
                p = doc.add_paragraph(style="List Bullet"); fmt(p)
                runs(p, x, cites, refs)
        elif kind == "figure":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=4)
            p.add_run().add_picture(str(OUT / "figures" / pl["file"]), width=Inches(3.0))
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
            runs(p, f"Fig. {refs[pl['label']]}. {pl['caption']}", cites, refs, size=8)
        elif kind == "table":
            if pl["wide"]:
                set_columns(doc.add_section(WD_SECTION.CONTINUOUS), 1)
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=6)
            runs(p, f"TABLE {refs[pl['label']]}", cites, refs, size=8)
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=3)
            runs(p, pl["caption"], cites, refs, size=8)
            rows = pl["rows"]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            total = 7.2 if pl["wide"] else 3.4
            weight = [min(max(max(len(plain(r[j], cites, refs).replace("*", "")) for r in rows), 5), 38)
                      for j in range(len(rows[0]))]
            widths = [Inches(total * w / sum(weight)) for w in weight]
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    c = t.cell(i, j)
                    c.width = widths[j]
                    c.text = ""
                    cp = c.paragraphs[0]
                    fmt(cp, WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
                    runs(cp, cell, cites, refs, size=7 if not pl["wide"] else 8, bold=(i == 0))
            p = doc.add_paragraph(); fmt(p, after=4)
            if pl["wide"]:
                set_columns(doc.add_section(WD_SECTION.CONTINUOUS), 2)
    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4)
    runs(p, "REFERENCES", cites, refs, size=10)
    for k, n in sorted(cites.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        runs(p, f"[{n}]\t{ALL_REFS[k]}", cites, refs, size=8)
    doc.save(path)


# ------------------------------------------------------------------ preview

def to_preview(title, blocks, cites, refs):
    out = [f"# {title}", "", "*Anonymous Authors (double-blind review)*", ""]
    sec_no, sub_no = 0, 0
    for kind, pl in blocks:
        if kind == "h1":
            if pl in ("Abstract", "Index Terms", "Acknowledgment"):
                out += [f"**{pl}**", ""]
            else:
                out += [f"## {ROMAN[sec_no]}. {pl}", ""]; sec_no += 1; sub_no = 0
        elif kind == "h2":
            out += [f"### {chr(65 + sub_no)}. {pl}", ""]; sub_no += 1
        elif kind == "p":
            out += [plain(pl, cites, refs), ""]
        elif kind == "list":
            out += [f"- {plain(x, cites, refs)}" for x in pl] + [""]
        elif kind == "figure":
            out += [f"![Fig. {refs[pl['label']]}](figures/{pl['file']})", "",
                    f"*Fig. {refs[pl['label']]}. {plain(pl['caption'], cites, refs)}*", ""]
        elif kind == "table":
            rows = pl["rows"]
            out += [f"*TABLE {refs[pl['label']]}. {plain(pl['caption'], cites, refs)}*", "",
                    "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
            out += ["| " + " | ".join(plain(c, cites, refs) for c in r) + " |" for r in rows[1:]] + [""]
    out += ["## References", ""]
    out += [f"[{n}] {ALL_REFS[k]}  " for k, n in sorted(cites.items(), key=lambda kv: kv[1])]
    return "\n".join(out) + "\n"


def main():
    title, blocks = parse((HERE / "paper_src.md").read_text(encoding="utf-8"))
    cites, refs = numbering(blocks)
    (OUT / "paper.tex").write_text(to_tex(title, blocks, cites), encoding="utf-8")
    to_docx(title, blocks, cites, refs, OUT / "paper.docx")
    preview = to_preview(title, blocks, cites, refs)
    (OUT / "PAPER_PREVIEW.md").write_text(preview, encoding="utf-8")
    body = [pl for kind, pl in blocks if kind in ("p", "list")]
    words = sum(len(plain(x if isinstance(x, str) else " ".join(x), cites, refs).split()) for x in body)
    print(f"wrote paper.tex, paper.docx, PAPER_PREVIEW.md | {len(cites)} references, "
          f"{sum(1 for k, _ in blocks if k == 'table')} tables, {sum(1 for k, _ in blocks if k == 'figure')} figure, "
          f"~{words} words of prose")


if __name__ == "__main__":
    main()
